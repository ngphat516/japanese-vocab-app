import csv
import io
from pathlib import Path
from typing import BinaryIO, Protocol
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from app.core.exceptions import AppError
from app.modules.imports.schemas import NormalizedRow


class Adapter(Protocol):
    def parse(self, file: BinaryIO) -> list[NormalizedRow]: ...


def read_table(file, required):
    name = getattr(file, "name", "upload.csv")
    payload = file.read(5 * 1024 * 1024 + 1)
    if len(payload) > 5 * 1024 * 1024:
        raise AppError(413, "IMPORT_FILE_TOO_LARGE", "File vượt quá 5 MB")
    suffix = Path(name).suffix.lower()
    if suffix == ".csv":
        try:
            text = payload.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise AppError(422, "IMPORT_ENCODING", "CSV phải dùng UTF-8") from None
        reader = csv.DictReader(io.StringIO(text, newline=""))
        fields = reader.fieldnames or []
        rows = []
        for row in reader:
            if None in row or any(v is None for v in row.values()):
                raise AppError(422, "IMPORT_COLUMNS", "Dòng CSV sai số cột")
            rows.append(dict(row))
            if len(rows) > 5000:
                raise AppError(422, "IMPORT_TOO_MANY_ROWS", "Tối đa 5.000 dòng")
    elif suffix == ".xlsx":
        try:
            with ZipFile(io.BytesIO(payload)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise AppError(422, "IMPORT_XLSX_TOO_LARGE", "XLSX giải nén vượt 50 MB")
            book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            try:
                values = book.active.iter_rows(values_only=True)
                fields = [str(x or "") for x in next(values)]
                rows = []
                for values_row in values:
                    if not any(x is not None for x in values_row):
                        continue
                    rows.append(
                        {
                            key: str(value) if value is not None else ""
                            for key, value in zip(fields, values_row, strict=True)
                        }
                    )
                    if len(rows) > 5000:
                        raise AppError(422, "IMPORT_TOO_MANY_ROWS", "Tối đa 5.000 dòng")
            finally:
                book.close()
        except (BadZipFile, StopIteration, KeyError, ValueError):
            raise AppError(422, "IMPORT_FORMAT", "XLSX không hợp lệ") from None
    else:
        raise AppError(422, "IMPORT_FORMAT", "Chỉ hỗ trợ CSV và XLSX")
    if len(fields) != len(set(fields)) or not required.issubset(set(fields)):
        raise AppError(
            422,
            "IMPORT_COLUMNS",
            "Thiếu hoặc trùng tên cột bắt buộc",
            {"required": sorted(required)},
        )
    if not rows:
        raise AppError(422, "IMPORT_EMPTY", "File không có dòng dữ liệu")
    return rows
